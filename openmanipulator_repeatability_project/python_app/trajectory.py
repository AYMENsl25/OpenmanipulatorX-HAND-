"""Trajectory records, Excel storage, and Cartesian replay validation."""

from __future__ import annotations

import math
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

import config
from experiment_frame import forward_kinematics_physical, inverse_kinematics_physical
from kinematics import JointAngles, MotorAngles, XYZ, motor_to_fk_angles


@dataclass(frozen=True)
class TrajectoryPoint:
    time_s: float
    motors: MotorAngles
    xyz: XYZ


@dataclass(frozen=True)
class XYZTrajectoryPoint:
    time_s: float
    xyz: XYZ


@dataclass(frozen=True)
class ValidationReport:
    total: int
    valid: int
    invalid: int
    maximum_error_mm: float
    average_error_mm: float
    errors: tuple[str, ...]


FULL_HEADERS = (
    "Index", "Time_s",
    "ID11_motor_deg", "ID12_motor_deg", "ID13_motor_deg", "ID14_motor_deg",
    "q1_rad", "q2_rad", "q3_rad", "q4_rad",
    "X_mm", "Y_mm", "Z_mm", "FK_raw_Z_mm", "Z_ground_clamped",
)
XYZ_HEADERS = ("Index", "Time_s", "X_mm", "Y_mm", "Z_mm")
MANUAL_HEADERS = (
    "Capture", "ID11_motor_deg", "ID12_motor_deg", "ID13_motor_deg", "ID14_motor_deg",
    "q1_deg", "q2_deg", "q3_deg", "q4_deg", "X_mm", "Y_mm", "Z_mm",
)


def _style_sheet(sheet) -> None:
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = fill
        cell.alignment = Alignment(horizontal="center")
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for column in sheet.columns:
        letter = column[0].column_letter
        width = max(len(str(cell.value)) if cell.value is not None else 0 for cell in column)
        sheet.column_dimensions[letter].width = min(max(width + 2, 12), 22)
    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            if isinstance(cell.value, float):
                cell.number_format = "0.000"


def _add_calibration_sheet(workbook: Workbook) -> None:
    sheet = workbook.create_sheet("Calibration")
    sheet.append(("Key", "Value"))
    sheet.append(("Kinematics_version", config.KINEMATICS_VERSION))
    sheet.append(("TCP", config.TCP_DESCRIPTION))
    sheet.append(("ID14_q4_zero_deg", 0.0))
    sheet.append(("Nominal_gripper_tip_extension_mm", config.NOMINAL_GRIPPER_TIP_EXTENSION))
    sheet.append(("Measured_TCP_ground_gap_correction_mm", config.MEASURED_TCP_GROUND_GAP_MM))
    sheet.append(("Gripper_tip_length_mm", config.L4_X))
    sheet.append(("Coordinate_frame", config.PHYSICAL_FRAME["origin"]))
    sheet.append(("Physical_axes", "+X forward; +Y right; +Z up"))
    sheet.append(("Physical_ground_z_mm", config.GROUND_Z_MM))
    sheet.append(("Teaching_negative_z_policy", "Replay Z is clamped to physical ground; raw FK Z is retained in Full_Trajectory"))
    sheet.append(("Experiment_config", str(config.EXPERIMENT_CONFIG_PATH)))
    _style_sheet(sheet)


def endpoint_points(points: list[TrajectoryPoint]) -> tuple[TrajectoryPoint, TrajectoryPoint]:
    if not points:
        raise ValueError("Trajectory is empty")
    return points[0], points[-1]


def xyz_endpoint_points(points: list[XYZTrajectoryPoint]) -> tuple[XYZTrajectoryPoint, XYZTrajectoryPoint]:
    if not points:
        raise ValueError("XYZ trajectory is empty")
    return points[0], points[-1]


def _validate_point(point: TrajectoryPoint) -> None:
    values = (point.time_s, *point.motors.as_tuple(), point.xyz.x, point.xyz.y, point.xyz.z)
    if not all(math.isfinite(value) for value in values):
        raise ValueError("Trajectory contains a non-finite value")


def ground_safe_teaching_xyz(motors: MotorAngles) -> tuple[XYZ, float, bool]:
    """Convert encoder FK to replay XYZ and clamp only values below ground.

    The unclamped FK Z is returned separately for diagnostics. X and Y and all
    motor/joint values remain untouched.
    """
    raw = forward_kinematics_physical(motors)
    was_clamped = raw.z < config.GROUND_Z_MM
    safe = XYZ(raw.x, raw.y, max(raw.z, config.GROUND_Z_MM))
    return safe, raw.z, was_clamped


def ground_safe_teaching_points(points: list[TrajectoryPoint]) -> list[TrajectoryPoint]:
    """Return teaching points whose replay XYZ never goes below ground."""
    return [
        TrajectoryPoint(point.time_s, point.motors, ground_safe_teaching_xyz(point.motors)[0])
        for point in points
    ]


def _timestamped_path(path: Path, stamp: str) -> Path:
    return path.with_name(f"{path.stem}_{stamp}{path.suffix}")


def _save_workbook_pair_with_lock_fallback(
    full: Workbook,
    xyz: Workbook,
    full_path: Path,
    xyz_path: Path,
) -> tuple[Path, Path]:
    """Save normal names, or a matched timestamped pair when Excel locks them."""
    try:
        full.save(full_path)
        xyz.save(xyz_path)
        return full_path, xyz_path
    except PermissionError:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        fallback_full = _timestamped_path(full_path, stamp)
        fallback_xyz = _timestamped_path(xyz_path, stamp)
        full.save(fallback_full)
        xyz.save(fallback_xyz)
        return fallback_full, fallback_xyz


def _save_workbook_with_lock_fallback(workbook: Workbook, path: Path) -> Path:
    try:
        workbook.save(path)
        return path
    except PermissionError:
        stamp = datetime.now().strftime("%Y%m%d_%H%M%S_%f")
        fallback = _timestamped_path(path, stamp)
        workbook.save(fallback)
        return fallback


def save_workbooks(points: list[TrajectoryPoint], directory: str | Path = ".") -> tuple[Path, Path]:
    if not points:
        raise ValueError("Cannot save an empty trajectory")
    output_dir = Path(directory)
    output_dir.mkdir(parents=True, exist_ok=True)
    full_path = output_dir / "trajectory_full.xlsx"
    xyz_path = output_dir / "trajectory_xyz.xlsx"

    full = Workbook()
    sheet = full.active
    sheet.title = "Full_Trajectory"
    sheet.append(FULL_HEADERS)
    for index, point in enumerate(points):
        _validate_point(point)
        joints = motor_to_fk_angles(point.motors)
        xyz_value, raw_z, was_clamped = ground_safe_teaching_xyz(point.motors)
        sheet.append((
            index,
            point.time_s,
            *point.motors.as_tuple(),
            *(math.radians(v) for v in joints.as_tuple()),
            xyz_value.x,
            xyz_value.y,
            xyz_value.z,
            raw_z,
            was_clamped,
        ))
    _style_sheet(sheet)
    _add_calibration_sheet(full)
    xyz = Workbook()
    sheet = xyz.active
    sheet.title = "XYZ_Trajectory"
    sheet.append(XYZ_HEADERS)
    for index, point in enumerate(points):
        xyz_value, _, _ = ground_safe_teaching_xyz(point.motors)
        sheet.append((index, point.time_s, xyz_value.x, xyz_value.y, xyz_value.z))
    _style_sheet(sheet)
    _add_calibration_sheet(xyz)
    return _save_workbook_pair_with_lock_fallback(full, xyz, full_path, xyz_path)


def save_manual_poses(points: list[TrajectoryPoint], directory: str | Path = ".") -> Path:
    if not points:
        raise ValueError("No manual poses have been captured")
    output_dir = Path(directory)
    output_dir.mkdir(parents=True, exist_ok=True)
    path = output_dir / "manual_poses.xlsx"
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Manual_Poses"
    sheet.append(MANUAL_HEADERS)
    for index, point in enumerate(points, start=1):
        _validate_point(point)
        joints = motor_to_fk_angles(point.motors)
        sheet.append((
            index,
            *point.motors.as_tuple(),
            *joints.as_tuple(),
            point.xyz.x,
            point.xyz.y,
            point.xyz.z,
        ))
    _style_sheet(sheet)
    _add_calibration_sheet(workbook)
    return _save_workbook_with_lock_fallback(workbook, path)


def load_xyz_workbook(path: str | Path) -> list[XYZTrajectoryPoint]:
    workbook = load_workbook(path, read_only=True, data_only=True)
    if "Calibration" not in workbook.sheetnames:
        workbook.close()
        raise ValueError(
            "This trajectory uses the old FK calibration. Record and save a new trajectory "
            "with the updated ID14/q4 and gripper-tip model."
        )
    calibration = workbook["Calibration"]
    saved_version = calibration["B2"].value
    if saved_version != config.KINEMATICS_VERSION:
        workbook.close()
        raise ValueError(
            f"Trajectory calibration {saved_version!r} does not match current "
            f"{config.KINEMATICS_VERSION!r}. Record a new trajectory."
        )
    sheet = workbook["XYZ_Trajectory"]
    headers = tuple(next(sheet.iter_rows(values_only=True)))
    if headers != XYZ_HEADERS:
        raise ValueError(f"Expected XYZ-only headers {XYZ_HEADERS}, received {headers}")
    points: list[XYZTrajectoryPoint] = []
    for row in sheet.iter_rows(min_row=2, values_only=True):
        if all(value is None for value in row):
            continue
        if len(row) != 5:
            raise ValueError("Malformed XYZ trajectory row")
        _, time_s, x, y, z = row
        points.append(XYZTrajectoryPoint(float(time_s), XYZ(float(x), float(y), float(z))))
    workbook.close()
    if not points:
        raise ValueError("XYZ trajectory is empty")
    return points


def validate_xyz_trajectory(points: list[XYZTrajectoryPoint], reference: JointAngles | None = None) -> ValidationReport:
    if not points:
        raise ValueError("XYZ trajectory is empty")
    errors: list[str] = []
    distances: list[float] = []
    current_reference = reference
    for index, target in enumerate(points):
        try:
            physical_result = inverse_kinematics_physical(
                target.xyz.x, target.xyz.y, target.xyz.z, reference=current_reference
            )
            current_reference = physical_result.ik.joint_angles
            error = physical_result.ik.position_error
            distances.append(error)
            if error > config.IK_POSITION_TOLERANCE_MM:
                errors.append(f"index={index} target=({target.xyz.x:.3f},{target.xyz.y:.3f},{target.xyz.z:.3f}) error={error:.3f} mm")
        except Exception as exc:
            errors.append(f"index={index} target=({target.xyz.x:.3f},{target.xyz.y:.3f},{target.xyz.z:.3f}) error={exc}")
    return ValidationReport(len(points), len(points) - len(errors), len(errors), max(distances, default=math.inf), sum(distances) / len(distances) if distances else math.inf, tuple(errors))
