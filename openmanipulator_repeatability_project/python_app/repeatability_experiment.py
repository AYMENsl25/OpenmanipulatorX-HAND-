"""Repeatability experiment records and recoverable Excel export.

The experiment compares the commanded physical point with FK reconstructed from
fresh encoder readings. XYZ comparisons, motor comparisons and diagnostics are
exported separately. Encoder FK is not independent physical metrology.
"""

from __future__ import annotations

from dataclasses import dataclass, field
from datetime import datetime, timezone
import math
import os
from pathlib import Path
from statistics import fmean, pstdev
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.chart import LineChart, BarChart, Reference

import config
from kinematics import JointAngles, MotorAngles, XYZ
from serial_controller import TelemetrySnapshot


def utc_timestamp() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="milliseconds").replace(
        "+00:00", "Z"
    )


def _mean(values: Iterable[float]) -> float:
    materialized = tuple(float(value) for value in values)
    return fmean(materialized)


def _std(values: Iterable[float]) -> float:
    materialized = tuple(float(value) for value in values)
    return pstdev(materialized) if len(materialized) > 1 else 0.0


def _vector_error(actual: XYZ, target: XYZ) -> tuple[float, float, float, float]:
    dx = actual.x - target.x
    dy = actual.y - target.y
    dz = actual.z - target.z
    return dx, dy, dz, math.sqrt(dx * dx + dy * dy + dz * dz)


@dataclass(frozen=True)
class RepeatabilityRunConfig:
    session_id: str
    selected_points: tuple[str, ...]
    repetitions: int
    sample_count: int
    sample_interval_ms: int
    pause_for_manual_measurement: bool = False
    speed_scale: float = 1.0
    started_utc: str = field(default_factory=utc_timestamp)


@dataclass(frozen=True)
class ManualMeasurement:
    measured_xyz: XYZ | None = None
    reported_distance_error_mm: float | None = None
    note: str = ""
    timestamp_utc: str = field(default_factory=utc_timestamp)


@dataclass
class PlannedPathPhase:
    cycle: int
    point_name: str
    phase: str
    physical_target: XYZ
    internal_target: XYZ
    planned_joints: JointAngles
    planned_motors: MotorAngles
    fk_check_physical: XYZ | None = None
    fk_error_mm: float | None = None


@dataclass
class TouchRecord:
    cycle: int
    point_name: str
    touch_index: int
    commanded_physical: XYZ
    commanded_internal: XYZ
    planned_joints: JointAngles
    planned_motors: MotorAngles
    samples: list[TelemetrySnapshot]
    planned_path: tuple[PlannedPathPhase, ...] = field(default_factory=tuple)
    recorded_utc: str = field(default_factory=utc_timestamp)
    manual: ManualMeasurement | None = None
    status: str = "AUTO_SAMPLED"

    def mean_physical_fk(self) -> XYZ:
        return XYZ(
            _mean(sample.state.physical_xyz.x for sample in self.samples),
            _mean(sample.state.physical_xyz.y for sample in self.samples),
            _mean(sample.state.physical_xyz.z for sample in self.samples),
        )

    def mean_internal_fk(self) -> XYZ:
        return XYZ(
            _mean(sample.state.internal_xyz.x for sample in self.samples),
            _mean(sample.state.internal_xyz.y for sample in self.samples),
            _mean(sample.state.internal_xyz.z for sample in self.samples),
        )

    def fk_std(self) -> XYZ:
        return XYZ(
            _std(sample.state.physical_xyz.x for sample in self.samples),
            _std(sample.state.physical_xyz.y for sample in self.samples),
            _std(sample.state.physical_xyz.z for sample in self.samples),
        )

    def fk_error(self) -> tuple[float, float, float, float]:
        return _vector_error(self.mean_physical_fk(), self.commanded_physical)

    def mean_raw_positions(self) -> tuple[float, float, float, float]:
        return tuple(
            _mean(sample.motors[motor_index].position_raw for sample in self.samples)
            for motor_index in range(4)
        )

    def mean_motor_angles(self) -> MotorAngles:
        return MotorAngles(
            *(
                _mean(sample.state.motors.as_tuple()[motor_index] for sample in self.samples)
                for motor_index in range(4)
            )
        )

    def mean_joint_angles(self) -> JointAngles:
        return JointAngles(
            *(
                _mean(sample.state.joints.as_tuple()[motor_index] for sample in self.samples)
                for motor_index in range(4)
            )
        )

    def manual_error(self) -> tuple[float, float, float, float] | None:
        if self.manual is None or self.manual.measured_xyz is None:
            return None
        return _vector_error(self.manual.measured_xyz, self.commanded_physical)


class RepeatabilityWorkbook:
    """Accumulate one run and rewrite a recoverable workbook after each touch."""

    def __init__(
        self,
        project_root: str | Path,
        run: RepeatabilityRunConfig,
        *,
        output_path: str | Path | None = None,
    ) -> None:
        self.run = run
        self.records: list[TouchRecord] = []
        self.events: list[tuple[str, str, str]] = []
        if output_path is None:
            self.run_dir = Path(project_root).resolve() / "experiment_results" / run.session_id
            self.run_dir.mkdir(parents=True, exist_ok=True)
            self.path = self.run_dir / "repeatability_point_results.xlsx"
            self.detail_path = self.run_dir / "repeatability_telemetry_details.xlsx"
        else:
            self.path = Path(output_path).resolve()
            self.run_dir = self.path.parent
            if not self.run_dir.is_dir():
                raise ValueError("output_path parent directory must already exist")
            self.detail_path = self.path.with_name(self.path.stem + "_details.xlsx")
        self.motor_path = self.path.with_name(self.path.stem + "_motor_angles.xlsx")

    def add_event(self, event: str, details: str = "") -> None:
        self.events.append((utc_timestamp(), event, details))

    def add_touch(self, record: TouchRecord) -> None:
        if not record.samples:
            raise ValueError("A touch record requires at least one telemetry sample")
        self.records.append(record)
        self.add_event(
            "touch_sampled",
            f"cycle={record.cycle}; point={record.point_name}; samples={len(record.samples)}",
        )
        self.save()

    def update_manual(self, record: TouchRecord, manual: ManualMeasurement | None) -> None:
        if record not in self.records:
            raise ValueError("Touch record does not belong to this run")
        record.manual = manual
        record.status = "MANUAL_RECORDED" if manual is not None else "MANUAL_SKIPPED"
        self.add_event(record.status.lower(), f"cycle={record.cycle}; point={record.point_name}")
        self.save()

    @staticmethod
    def _style_sheet(sheet) -> None:
        header_fill = PatternFill("solid", fgColor="1F4E78")
        for cell in sheet[1]:
            cell.font = Font(color="FFFFFF", bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for column in sheet.columns:
            values = [str(cell.value) if cell.value is not None else "" for cell in column[:100]]
            width = min(24, max(9, max((len(value) for value in values), default=0) + 2))
            sheet.column_dimensions[column[0].column_letter].width = width

    @staticmethod
    def _append_xyz(row: list, value: XYZ | None) -> None:
        row.extend((value.x, value.y, value.z) if value is not None else (None, None, None))

    def _summary_headers(self) -> list[str]:
        return [
            "session_id", "cycle", "point", "touch_index", "recorded_utc", "status",
            "target_x_mm", "target_y_mm", "target_z_mm",
            "fk_mean_x_mm", "fk_mean_y_mm", "fk_mean_z_mm",
            "fk_error_x_mm", "fk_error_y_mm", "fk_error_z_mm", "fk_error_norm_mm",
            "fk_std_x_mm", "fk_std_y_mm", "fk_std_z_mm", "fk_noise_norm_mm",
            "ik_id11_motor_deg", "ik_id12_motor_deg", "ik_id13_motor_deg", "ik_id14_motor_deg",
            "actual_id11_motor_deg", "actual_id12_motor_deg",
            "actual_id13_motor_deg", "actual_id14_motor_deg",
            "actual_id11_position_raw_mean", "actual_id12_position_raw_mean",
            "actual_id13_position_raw_mean", "actual_id14_position_raw_mean",
            "sample_count",
            "manual_x_mm", "manual_y_mm", "manual_z_mm",
            "manual_error_x_mm", "manual_error_y_mm", "manual_error_z_mm",
            "manual_error_norm_mm", "manual_reported_distance_error_mm",
            "manual_note", "manual_timestamp_utc",
        ]

    def _summary_row(self, record: TouchRecord) -> list:
        mean_fk = record.mean_physical_fk()
        std_fk = record.fk_std()
        fk_error = record.fk_error()
        actual_motors = record.mean_motor_angles()
        manual_xyz = record.manual.measured_xyz if record.manual else None
        manual_error = record.manual_error()
        row = [
            self.run.session_id, record.cycle, record.point_name, record.touch_index,
            record.recorded_utc, record.status,
        ]
        self._append_xyz(row, record.commanded_physical)
        self._append_xyz(row, mean_fk)
        row.extend(fk_error)
        self._append_xyz(row, std_fk)
        row.extend((math.sqrt(std_fk.x ** 2 + std_fk.y ** 2 + std_fk.z ** 2),))
        row.extend(record.planned_motors.as_tuple())
        row.extend(actual_motors.as_tuple())
        row.extend(record.mean_raw_positions())
        row.append(len(record.samples))
        self._append_xyz(row, manual_xyz)
        row.extend(manual_error if manual_error is not None else (None, None, None, None))
        row.extend(
            (
                record.manual.reported_distance_error_mm if record.manual else None,
                record.manual.note if record.manual else "",
                record.manual.timestamp_utc if record.manual else "",
            )
        )

        return row

    def _path_headers(self) -> list[str]:
        return [
            "session_id", "cycle", "point", "phase",
            "target_x_mm", "target_y_mm", "target_z_mm",
            "internal_target_x_mm", "internal_target_y_mm", "internal_target_z_mm",
            "planned_q1_deg", "planned_q2_deg", "planned_q3_deg", "planned_q4_deg",
            "planned_id11_motor_deg", "planned_id12_motor_deg",
            "planned_id13_motor_deg", "planned_id14_motor_deg",
            "fk_check_x_mm", "fk_check_y_mm", "fk_check_z_mm", "ik_fk_error_mm",
        ]

    def _path_row(self, phase: PlannedPathPhase) -> list:
        row = [
            self.run.session_id, phase.cycle, phase.point_name, phase.phase,
        ]
        self._append_xyz(row, phase.physical_target)
        self._append_xyz(row, phase.internal_target)
        row.extend(phase.planned_joints.as_tuple())
        row.extend(phase.planned_motors.as_tuple())
        self._append_xyz(row, phase.fk_check_physical)
        row.append(phase.fk_error_mm)
        return row

    @staticmethod
    def _bitwise_or(values: Iterable[int]) -> int:
        result = 0
        for value in values:
            result |= int(value)
        return result

    def _sample_headers(self) -> list[str]:
        headers = [
            "session_id", "cycle", "point", "touch_index", "sample_index", "timestamp_utc",
            "command_x_mm", "command_y_mm", "command_z_mm",
            "fk_x_mm", "fk_y_mm", "fk_z_mm",
            "fk_error_x_mm", "fk_error_y_mm", "fk_error_z_mm", "fk_error_norm_mm",
        ]
        for motor_id in (11, 12, 13, 14):
            prefix = f"id{motor_id}"
            headers.extend(
                (
                    f"{prefix}_position_raw", f"{prefix}_motor_deg", f"{prefix}_q_deg",
                    f"{prefix}_velocity_raw", f"{prefix}_velocity_rpm",
                    f"{prefix}_current_raw", f"{prefix}_current_ma",
                    f"{prefix}_pwm_raw", f"{prefix}_pwm_percent",
                    f"{prefix}_voltage_raw", f"{prefix}_voltage_v",
                    f"{prefix}_temperature_c", f"{prefix}_hardware_error",
                    f"{prefix}_moving", f"{prefix}_moving_status",
                )
            )
        return headers

    def _sample_row(self, record: TouchRecord, sample_index: int, sample: TelemetrySnapshot) -> list:
        state = sample.state
        error = _vector_error(state.physical_xyz, record.commanded_physical)
        row = [
            self.run.session_id, record.cycle, record.point_name, record.touch_index,
            sample_index, sample.timestamp_utc,
            record.commanded_physical.x, record.commanded_physical.y, record.commanded_physical.z,
            state.physical_xyz.x, state.physical_xyz.y, state.physical_xyz.z,
            *error,
        ]
        for index, motor in enumerate(sample.motors):
            row.extend(
                (
                    motor.position_raw, state.motors.as_tuple()[index], state.joints.as_tuple()[index],
                    motor.velocity_raw, motor.velocity_rpm,
                    motor.current_raw, motor.current_ma,
                    motor.pwm_raw, motor.pwm_percent,
                    motor.voltage_raw, motor.voltage_v,
                    motor.temperature_c, motor.hardware_error, motor.moving, motor.moving_status,
                )
            )
        return row

    def _build_point_workbook(self) -> Workbook:
        workbook = Workbook()
        summary = workbook.active
        summary.title = "Point Results"
        all_headers = self._summary_headers()
        keep = [i for i, name in enumerate(all_headers) if not name.startswith(("manual_", "ik_id", "actual_id"))]
        summary.append([all_headers[i] for i in keep])
        for record in self.records:
            row = self._summary_row(record)
            summary.append([row[i] for i in keep])

        plots = workbook.create_sheet("Error Plots")
        plots.append(("Encoder FK comparisons only; no independent physical measurement",))
        if self.records:
            for title, columns, anchor in (
                ("Signed XYZ error by touch", (13, 15), "A3"),
                ("Distance error and sample noise by touch", (16, 16), "A20"),
            ):
                chart = LineChart()
                chart.title = title
                chart.y_axis.title = "mm"
                chart.x_axis.title = "Touch index"
                chart.add_data(Reference(summary, min_col=columns[0], max_col=columns[1], min_row=1, max_row=summary.max_row), titles_from_data=True)
                if columns == (16, 16):
                    chart.add_data(Reference(summary, min_col=20, min_row=1, max_row=summary.max_row), titles_from_data=True)
                chart.set_categories(Reference(summary, min_col=4, min_row=2, max_row=summary.max_row))
                chart.width, chart.height = 25, 9
                plots.add_chart(chart, anchor)
            aggregate = workbook.create_sheet("Point Error Summary")
            aggregate.append(("point", "mean_distance_error_mm", "max_distance_error_mm", "repeatability_x_std_mm", "repeatability_y_std_mm", "repeatability_z_std_mm", "touch_count"))
            for name in self.run.selected_points:
                records = [r for r in self.records if r.point_name == name]
                if not records:
                    continue
                errors = [r.fk_error()[3] for r in records]
                positions = [r.mean_physical_fk() for r in records]
                aggregate.append((name, fmean(errors), max(errors), *(_std(getattr(p, axis) for p in positions) for axis in ("x", "y", "z")), len(records)))
            chart = BarChart()
            chart.title = "Mean and maximum encoder FK error per point"
            chart.y_axis.title = "mm"
            chart.add_data(Reference(aggregate, min_col=2, max_col=3, min_row=1, max_row=aggregate.max_row), titles_from_data=True)
            chart.set_categories(Reference(aggregate, min_col=1, min_row=2, max_row=aggregate.max_row))
            chart.width, chart.height = 25, 9
            plots.add_chart(chart, "A37")

        path = workbook.create_sheet("Planned Point Path")
        path.append(self._path_headers())
        for record in self.records:
            for phase in record.planned_path:
                path.append(self._path_row(phase))

        run_config = workbook.create_sheet("Run Config")
        run_config.append(("field", "value"))
        run_config_rows = (
            ("session_id", self.run.session_id),
            ("started_utc", self.run.started_utc),
            ("selected_points", ", ".join(self.run.selected_points)),
            ("repetitions", self.run.repetitions),
            ("touch_sample_count", self.run.sample_count),
            ("sample_interval_ms", self.run.sample_interval_ms),
            ("speed_scale", self.run.speed_scale),
            ("motor_arrival_tolerance_deg", config.POST_MOVE_TOLERANCE_DEGREES),
            ("motor_tracking_warning_deg", config.MOTOR_TRACKING_WARNING_DEGREES),
            ("measurement_mode", "automatic encoder FK"),
            ("kinematics_version", config.KINEMATICS_VERSION),
            ("tcp_definition", config.TCP_DESCRIPTION),
            ("coordinate_units", "millimetres"),
            ("physical_frame", "ID11 center; +X forward; +Y right; +Z up"),
            ("work_xyz_mm", str(config.WORK_XYZ_MM)),
            ("approach_clearance_mm", config.APPROACH_CLEARANCE_MM),
            ("retract_clearance_mm", config.RETRACT_CLEARANCE_MM),
        )
        for item in run_config_rows:
            run_config.append(item)

        for sheet in workbook.worksheets:
            self._style_sheet(sheet)
        return workbook

    def _build_motor_workbook(self) -> Workbook:
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Motor Angle Comparison"
        headers = ["session_id", "cycle", "point", "touch_index"]
        for motor_id in (11, 12, 13, 14):
            headers.extend((f"id{motor_id}_ik_motor_deg", f"id{motor_id}_actual_motor_deg", f"id{motor_id}_motor_error_deg", f"id{motor_id}_ik_q_deg", f"id{motor_id}_actual_q_deg", f"id{motor_id}_q_error_deg", f"id{motor_id}_raw_mean"))
        sheet.append(headers)
        for record in self.records:
            row = [self.run.session_id, record.cycle, record.point_name, record.touch_index]
            actual_m = record.mean_motor_angles().as_tuple()
            actual_q = record.mean_joint_angles().as_tuple()
            for i in range(4):
                planned_m = record.planned_motors.as_tuple()[i]
                planned_q = record.planned_joints.as_tuple()[i]
                # Wrapped motor display differences; calibrated q stays continuous.
                motor_error = (actual_m[i] - planned_m + 180) % 360 - 180
                row.extend((planned_m, actual_m[i], motor_error, planned_q, actual_q[i], actual_q[i] - planned_q, record.mean_raw_positions()[i]))
            sheet.append(row)
        self._style_sheet(sheet)
        if self.records:
            chart = LineChart()
            chart.title = "Calibrated joint tracking error by touch"
            chart.y_axis.title = "degrees"
            for column in (10, 17, 24, 31):
                chart.add_data(Reference(sheet, min_col=column, min_row=1, max_row=sheet.max_row), titles_from_data=True)
            chart.set_categories(Reference(sheet, min_col=4, min_row=2, max_row=sheet.max_row))
            plots = workbook.create_sheet("Motor Error Plot")
            plots.add_chart(chart, "A1")
        return workbook

    def _build_detail_workbook(self) -> Workbook:
        workbook = Workbook()
        samples = workbook.active
        samples.title = "Telemetry Samples"
        samples.append(self._sample_headers())
        for record in self.records:
            for sample_index, sample in enumerate(record.samples, start=1):
                samples.append(self._sample_row(record, sample_index, sample))

        events = workbook.create_sheet("Run Events")
        events.append(("timestamp_utc", "event", "details"))
        for event in self.events:
            events.append(event)

        for sheet in workbook.worksheets:
            self._style_sheet(sheet)
        return workbook

    def _save_one(self, workbook: Workbook, path: Path, recovery_prefix: str) -> Path:
        temporary = path.with_name(path.stem + ".saving.xlsx")
        try:
            workbook.save(temporary)
            os.replace(temporary, path)
            return path
        except PermissionError:
            temporary.unlink(missing_ok=True)
            timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
            recovery = self.run_dir / f"{recovery_prefix}_{timestamp}.xlsx"
            workbook.save(recovery)
            return recovery
        finally:
            workbook.close()

    def save(self) -> Path:
        self.path = self._save_one(
            self._build_point_workbook(), self.path, "repeatability_point_results"
        )
        self.detail_path = self._save_one(
            self._build_detail_workbook(), self.detail_path, "repeatability_telemetry_details"
        )
        self.motor_path = self._save_one(self._build_motor_workbook(), self.motor_path, "repeatability_motor_angles")
        return self.path


__all__ = [
    "ManualMeasurement",
    "PlannedPathPhase",
    "RepeatabilityRunConfig",
    "RepeatabilityWorkbook",
    "TouchRecord",
]
