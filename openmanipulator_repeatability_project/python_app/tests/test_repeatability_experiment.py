from __future__ import annotations

from pathlib import Path
import sys
import unittest

from openpyxl import load_workbook

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from kinematics import JointAngles, MotorAngles, XYZ
from repeatability_experiment import (
    ManualMeasurement,
    PlannedPathPhase,
    RepeatabilityRunConfig,
    RepeatabilityWorkbook,
    TouchRecord,
)
from serial_controller import ControllerError, MotorTelemetry, OpenCRController, RobotState, TelemetrySnapshot


def telemetry_snapshot(offset: int = 0) -> TelemetrySnapshot:
    motors = MotorAngles(168.486, 179.824, 355.869, 82.881)
    joints = JointAngles(0.0, 0.0, 0.0, 82.881)
    physical = XYZ(230.0 + offset * 0.1, 70.0, offset * 0.05)
    state = RobotState(
        timestamp_utc="2026-09-10T00:00:00.000Z",
        raw_positions=(1917 + offset, 2046, 4049, 943),
        motors=motors,
        joints=joints,
        physical_xyz=physical,
        internal_xyz=XYZ(physical.x, -physical.y, physical.z),
    )
    diagnostics = tuple(
        MotorTelemetry(
            motor_id=motor_id,
            position_raw=raw,
            velocity_raw=offset,
            current_raw=10 + offset,
            pwm_raw=20 + offset,
            voltage_raw=120,
            temperature_c=30 + offset,
            hardware_error=0,
            moving=0,
            moving_status=0,
        )
        for motor_id, raw in zip((11, 12, 13, 14), state.raw_positions, strict=True)
    )
    return TelemetrySnapshot(state.timestamp_utc, state, diagnostics)


class RepeatabilityExperimentTests(unittest.TestCase):
    def test_collection_arrival_tolerance_boundary(self) -> None:
        controller = OpenCRController()
        target = MotorAngles(100, 180, 350, 80)
        for error in (0.791, 9.0, 10.0):
            controller.read_motor_angles = lambda: MotorAngles(100, 180 + error, 350, 80)
            controller._verify_motor_targets(target)
            self.assertAlmostEqual(controller.last_motor_tracking_errors[1], error)
        controller.read_motor_angles = lambda: MotorAngles(100, 190.1, 350, 80)
        with self.assertRaises(ControllerError):
            controller._verify_motor_targets(target)

    def test_motion_speed_protocol_and_validation(self) -> None:
        controller = OpenCRController()
        commands = []
        def command(value):
            commands.append(value)
            return "OK,SPEED"
        controller._command = command
        controller.set_motion_speed(0.5)
        self.assertEqual(commands, ["SET_SPEED,0.500"])
        for value in (0.0, 1.1, float("nan")):
            with self.assertRaises(ValueError):
                controller.set_motion_speed(value)
        self.assertEqual(len(commands), 1)

    def test_telemetry_protocol_parser(self) -> None:
        groups = []
        for motor_id, raw in zip((11, 12, 13, 14), (1917, 2046, 4049, 943), strict=True):
            groups.extend((motor_id, raw, -2, 10, 20, 120, 31, 0, 0, 0))
        response = "TELEMETRY_V1," + ",".join(str(value) for value in groups)
        controller = OpenCRController()
        controller._command = lambda command: response  # type: ignore[method-assign]

        snapshot = controller.read_telemetry()

        self.assertEqual(tuple(m.motor_id for m in snapshot.motors), (11, 12, 13, 14))
        self.assertEqual(snapshot.state.raw_positions, (1917, 2046, 4049, 943))
        self.assertAlmostEqual(snapshot.motors[0].velocity_rpm, -0.458)
        self.assertAlmostEqual(snapshot.motors[0].current_ma, 26.9)
        self.assertAlmostEqual(snapshot.motors[0].voltage_v, 12.0)

    def test_workbooks_separate_point_results_from_telemetry_details(self) -> None:
        output_path = Path(__file__).resolve().parent / "repeatability_test_output.xlsx"
        details_path = output_path.with_name("repeatability_test_output_details.xlsx")
        try:
            run = RepeatabilityRunConfig(
                session_id="test-repeatability",
                selected_points=("P09",),
                repetitions=2,
                sample_count=2,
                sample_interval_ms=50,
                pause_for_manual_measurement=True,
            )
            writer = RepeatabilityWorkbook(
                Path(__file__).resolve().parents[2], run, output_path=output_path
            )
            record = TouchRecord(
                cycle=1,
                point_name="P09",
                touch_index=1,
                commanded_physical=XYZ(230.0, 70.0, 0.0),
                commanded_internal=XYZ(230.0, -70.0, 0.0),
                planned_joints=JointAngles(0.0, 0.0, 0.0, 0.0),
                planned_motors=MotorAngles(168.486, 179.824, 355.869, 82.881),
                samples=[telemetry_snapshot(0), telemetry_snapshot(1)],
                planned_path=(
                    PlannedPathPhase(
                        cycle=1,
                        point_name="P09",
                        phase="TOUCH",
                        physical_target=XYZ(230.0, 70.0, 0.0),
                        internal_target=XYZ(230.0, -70.0, 0.0),
                        planned_joints=JointAngles(0.0, 0.0, 0.0, 0.0),
                        planned_motors=MotorAngles(168.486, 179.824, 355.869, 82.881),
                        fk_check_physical=XYZ(230.0, 70.0, 0.0),
                        fk_error_mm=0.0,
                    ),
                ),
            )
            writer.add_touch(record)

            self.assertTrue(writer.path.exists())
            self.assertTrue(writer.detail_path.exists())
            workbook = load_workbook(writer.path, data_only=False, read_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    ["Point Results", "Error Plots", "Point Error Summary", "Planned Point Path", "Run Config"],
                )
                summary = workbook["Point Results"]
                headers = [cell.value for cell in summary[1]]
                row = [cell.value for cell in summary[2]]
                self.assertEqual(row[headers.index("point")], "P09")
                self.assertAlmostEqual(row[headers.index("target_x_mm")], 230.0)
                self.assertAlmostEqual(row[headers.index("fk_actual_x_mm")], 230.05)
                self.assertNotIn("fk_mean_x_mm", headers)
                self.assertFalse(any(h.startswith("manual_") for h in headers))
                self.assertAlmostEqual(row[headers.index("fk_error_norm_mm")], (0.05**2 + 0.025**2)**0.5)
                self.assertEqual(workbook["Planned Point Path"].max_row, 2)
                self.assertIs(record.calculated_result(), record.calculated_result())
            finally:
                workbook.close()
            details = load_workbook(writer.detail_path, data_only=False, read_only=True)
            try:
                self.assertEqual(details.sheetnames, ["Telemetry Samples", "Run Events"])
                self.assertEqual(details["Telemetry Samples"].max_row, 3)
                self.assertGreater(details["Run Events"].max_row, 1)
            finally:
                details.close()
            motors = load_workbook(writer.motor_path, read_only=False)
            try:
                sheet = motors["Motor Angle Comparison"]
                data = dict(zip([c.value for c in sheet[1]], [c.value for c in sheet[2]]))
                self.assertAlmostEqual(data["id11_actual_motor_deg"], 168.486)
                self.assertAlmostEqual(data["id11_raw_mean"], 1917.5)
                self.assertEqual(len(motors["Motor Error Plot"]._charts), 1)
            finally:
                motors.close()
        finally:
            output_path.unlink(missing_ok=True)
            details_path.unlink(missing_ok=True)
            output_path.with_name(output_path.stem + "_motor_angles.xlsx").unlink(missing_ok=True)
            output_path.with_name("repeatability_test_output.saving.xlsx").unlink(missing_ok=True)
            details_path.with_name("repeatability_test_output_details.saving.xlsx").unlink(
                missing_ok=True
            )


if __name__ == "__main__":
    unittest.main()
