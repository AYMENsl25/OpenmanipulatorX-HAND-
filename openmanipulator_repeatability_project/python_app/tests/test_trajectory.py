import sys
import unittest
from openpyxl import load_workbook
from pathlib import Path

sys.path.insert(0, str(Path(__file__).parents[1]))

from experiment_frame import validate_physical_target
from kinematics import KinematicsError, MotorAngles, XYZ
from trajectory import TrajectoryPoint, XYZTrajectoryPoint, endpoint_points, ground_safe_teaching_points, ground_safe_teaching_xyz, save_workbooks, xyz_endpoint_points


def test_endpoint_points_preserve_first_and_last_teach_samples():
    first = TrajectoryPoint(0.0, MotorAngles(351, 1, 1, 0), XYZ(1, 2, 3))
    last = TrajectoryPoint(1.0, MotorAngles(350, 2, 3, 4), XYZ(4, 5, 6))
    start, end = endpoint_points([first, last])
    assert start == first
    assert end == last


def test_xyz_endpoint_points_work_after_loading_xyz_workbook():
    first = XYZTrajectoryPoint(0.0, XYZ(10, 20, 30))
    last = XYZTrajectoryPoint(1.0, XYZ(40, 50, 60))
    assert xyz_endpoint_points([first, last]) == (first, last)


def test_measured_workspace_rejects_previous_outlier():
    try:
        validate_physical_target(XYZ(74.563, 389.0, 359.303))
    except KinematicsError:
        return
    raise AssertionError("Measured workspace should reject Y=389 mm")


def test_teaching_ground_clamp_preserves_raw_fk_for_diagnostics():
    motors = MotorAngles(169.014, 209.971, 354.990, 42.100)
    safe, raw_z, was_clamped = ground_safe_teaching_xyz(motors)
    assert was_clamped
    assert raw_z < 0.0
    assert safe.z == 0.0


class TeachingGroundClampTests(unittest.TestCase):
    def test_negative_fk_z_becomes_zero_in_replay_point(self):
        source = TrajectoryPoint(
            1.0,
            MotorAngles(169.014, 209.971, 354.990, 42.100),
            XYZ(255.901, -2.357, -35.871),
        )
        safe = ground_safe_teaching_points([source])[0]
        self.assertEqual(safe.time_s, source.time_s)
        self.assertEqual(safe.motors, source.motors)
        self.assertEqual(safe.xyz.z, 0.0)

    def test_saved_replay_clamps_z_and_full_file_keeps_raw_z(self):
        point = TrajectoryPoint(
            1.0,
            MotorAngles(169.014, 209.971, 354.990, 42.100),
            XYZ(255.901, -2.357, -35.871),
        )
        directory = Path(__file__).resolve().parent / "trajectory_ground_test"
        directory.mkdir(exist_ok=True)
        full_path = directory / "trajectory_full.xlsx"
        xyz_path = directory / "trajectory_xyz.xlsx"
        try:
            full_path, xyz_path = save_workbooks([point], directory)
            xyz_book = load_workbook(xyz_path, read_only=True, data_only=True)
            self.assertEqual(xyz_book["XYZ_Trajectory"]["E2"].value, 0.0)
            xyz_book.close()
            full_book = load_workbook(full_path, read_only=True, data_only=True)
            headers = [cell.value for cell in full_book["Full_Trajectory"][1]]
            values = [cell.value for cell in full_book["Full_Trajectory"][2]]
            row = dict(zip(headers, values))
            self.assertEqual(row["Z_mm"], 0.0)
            self.assertLess(row["FK_raw_Z_mm"], 0.0)
            self.assertTrue(row["Z_ground_clamped"])
            full_book.close()
        finally:
            full_path.unlink(missing_ok=True)
            xyz_path.unlink(missing_ok=True)
