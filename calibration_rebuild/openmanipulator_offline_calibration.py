"""Offline OpenMANIPULATOR-X motor calibration, FK and position-only IK.

This program never opens a serial port and cannot move the robot.  It fits only
motor-to-joint offsets/signs and a constrained robot-to-work-frame transform;
the official link geometry remains fixed.
"""

from __future__ import annotations

import argparse
from dataclasses import dataclass
from itertools import product
import math
from pathlib import Path
from typing import Iterable, Sequence

import numpy as np
import pandas as pd
from scipy.optimize import least_squares


# ----------------------------- configuration -----------------------------
BASE_X_MM = 12.0
BASE_Z_MM = 17.0
J2_Z_MM = 59.5
J3_X_MM = 24.0
J3_Z_MM = 128.0
J4_X_MM = 124.0
TCP_X_MM = 126.0
RAW_PER_REV = 4096.0
DEG_PER_REV = 360.0

ROBUST_LOSS = "soft_l1"
ROBUST_F_SCALE_MM = 20.0
CALIBRATION_MULTISTARTS = 8
LOO_MULTISTARTS = 3
IK_MAX_ITERATIONS = 300
IK_TOLERANCE_MM = 0.05
IK_DAMPING = 2.0
IK_MAX_STEP_DEG = 8.0
IK_FINITE_DIFFERENCE_RAD = 1.0e-5
PHYSICAL_MEAN_ERROR_LIMIT_MM = 10.0
PHYSICAL_MAX_ERROR_LIMIT_MM = 20.0
PHYSICAL_LOO_ERROR_LIMIT_MM = 30.0

# Nominal model limits; these are not substitutes for measured safe RAW limits.
JOINT_LIMITS_DEG = np.array([
    [-180.0, 180.0],
    [math.degrees(-2.05), 90.0],
    [-90.0, math.degrees(1.53)],
    [math.degrees(-1.8), math.degrees(2.0)],
])

OUTPUT_FILENAME = "openmanipulator_fk_ik_validation.xlsx"

POSES = [
    dict(Pose="Pose 1", RAW11=1913, RAW12=2071, RAW13=4052, RAW14=922, RAW15=1209,
         MotorDeg11=168.135, MotorDeg12=182.021, MotorDeg13=356.133, MotorDeg14=81.035, MotorDeg15=106.260,
         Measured_X_mm=175.0, Measured_Y_mm=0.0, Measured_Z_mm=35.0),
    dict(Pose="Pose 2", RAW11=1655, RAW12=2830, RAW13=3047, RAW14=1190, RAW15=1210,
         MotorDeg11=145.459, MotorDeg12=248.730, MotorDeg13=267.803, MotorDeg14=104.590, MotorDeg15=106.348,
         Measured_X_mm=240.0, Measured_Y_mm=150.0, Measured_Z_mm=0.0),
    dict(Pose="Pose 3", RAW11=996, RAW12=2086, RAW13=4357, RAW14=752, RAW15=1209,
         MotorDeg11=87.539, MotorDeg12=183.340, MotorDeg13=382.939, MotorDeg14=66.094, MotorDeg15=106.260,
         Measured_X_mm=0.0, Measured_Y_mm=155.0, Measured_Z_mm=0.0),
    dict(Pose="Pose 4", RAW11=3012, RAW12=2037, RAW13=4442, RAW14=769, RAW15=2664,
         MotorDeg11=264.727, MotorDeg12=179.033, MotorDeg13=390.410, MotorDeg14=67.588, MotorDeg15=234.141,
         Measured_X_mm=0.0, Measured_Y_mm=-145.0, Measured_Z_mm=0.0),
    dict(Pose="Pose 5", RAW11=1940, RAW12=2077, RAW13=4067, RAW14=109, RAW15=2663,
         MotorDeg11=170.508, MotorDeg12=182.549, MotorDeg13=357.451, MotorDeg14=9.580, MotorDeg15=234.053,
         Measured_X_mm=290.0, Measured_Y_mm=0.0, Measured_Z_mm=210.0),
    dict(Pose="Pose 6", RAW11=2267, RAW12=2545, RAW13=3696, RAW14=608, RAW15=2663,
         MotorDeg11=199.248, MotorDeg12=223.682, MotorDeg13=324.844, MotorDeg14=53.438, MotorDeg15=234.053,
         Measured_X_mm=280.0, Measured_Y_mm=-140.0, Measured_Z_mm=0.0),
]


@dataclass(frozen=True)
class CalibrationFit:
    signs: np.ndarray
    zeros_deg: np.ndarray
    translation_mm: np.ndarray
    yaw_deg: float
    robust_cost: float
    residual_vector: np.ndarray
    success: bool


@dataclass(frozen=True)
class IKResult:
    converged: bool
    q_deg: np.ndarray
    reconstructed_work_mm: np.ndarray
    error_vector_mm: np.ndarray
    error_mm: float
    iterations: int
    reason: str


def motor_deg_to_q(motor_deg: np.ndarray, signs: np.ndarray, zeros_deg: np.ndarray) -> np.ndarray:
    """Convert continuous motor degrees to mathematical joint degrees."""
    return signs * (motor_deg - zeros_deg)


def q_to_motor_deg(q_deg: np.ndarray, signs: np.ndarray, zeros_deg: np.ndarray) -> np.ndarray:
    """Inverse calibration. No modulo wrapping is performed."""
    return zeros_deg + q_deg / signs


def motor_deg_to_raw(motor_deg: np.ndarray) -> np.ndarray:
    """Convert continuous motor degrees to continuous RAW counts."""
    return motor_deg * RAW_PER_REV / DEG_PER_REV


def forward_kinematics_robot(q_deg: Sequence[float]) -> np.ndarray:
    """Explicit official transform chain; returns robot-frame XYZ in mm."""
    q1, q2, q3, q4 = np.radians(np.asarray(q_deg, dtype=float))
    q23 = q2 + q3
    q234 = q23 + q4

    x23 = math.cos(q2) * J3_X_MM + math.sin(q2) * J3_Z_MM
    z23 = -math.sin(q2) * J3_X_MM + math.cos(q2) * J3_Z_MM
    x34 = math.cos(q23) * J4_X_MM
    z34 = -math.sin(q23) * J4_X_MM
    x45 = math.cos(q234) * TCP_X_MM
    z45 = -math.sin(q234) * TCP_X_MM

    radial = x23 + x34 + x45
    vertical = J2_Z_MM + z23 + z34 + z45
    return np.array([
        BASE_X_MM + math.cos(q1) * radial,
        math.sin(q1) * radial,
        BASE_Z_MM + vertical,
    ])


def tool_pitch_deg(q_deg: Sequence[float]) -> float:
    q = np.asarray(q_deg, dtype=float)
    return float(q[1] + q[2] + q[3])


def robot_to_work(robot_xyz: np.ndarray, translation_mm: np.ndarray, yaw_deg: float) -> np.ndarray:
    yaw = math.radians(yaw_deg)
    c, s = math.cos(yaw), math.sin(yaw)
    rotation = np.array([[c, -s, 0.0], [s, c, 0.0], [0.0, 0.0, 1.0]])
    return rotation @ np.asarray(robot_xyz, dtype=float) + translation_mm


def work_to_robot(work_xyz: np.ndarray, translation_mm: np.ndarray, yaw_deg: float) -> np.ndarray:
    yaw = math.radians(yaw_deg)
    c, s = math.cos(yaw), math.sin(yaw)
    rotation_transpose = np.array([[c, s, 0.0], [-s, c, 0.0], [0.0, 0.0, 1.0]])
    return rotation_transpose @ (np.asarray(work_xyz, dtype=float) - translation_mm)


def forward_kinematics_work(q_deg: Sequence[float], fit: CalibrationFit) -> np.ndarray:
    return robot_to_work(forward_kinematics_robot(q_deg), fit.translation_mm, fit.yaw_deg)


def _arrays(rows: Sequence[dict]) -> tuple[np.ndarray, np.ndarray]:
    motor = np.array([[r[f"MotorDeg{i}"] for i in range(11, 15)] for r in rows], dtype=float)
    measured = np.array([[r["Measured_X_mm"], r["Measured_Y_mm"], r["Measured_Z_mm"]] for r in rows], dtype=float)
    return motor, measured


def calibration_residual(parameters: np.ndarray, signs: np.ndarray,
                         motor_deg: np.ndarray, measured_xyz: np.ndarray) -> np.ndarray:
    zeros = parameters[:4]
    translation = parameters[4:7]
    yaw = parameters[7]
    residuals = []
    for motors, measured in zip(motor_deg, measured_xyz, strict=True):
        q_deg = motor_deg_to_q(motors, signs, zeros)
        predicted = robot_to_work(forward_kinematics_robot(q_deg), translation, yaw)
        residuals.extend(predicted - measured)
    return np.asarray(residuals)


def _fit_one_sign(signs: np.ndarray, rows: Sequence[dict], starts: int,
                  rng: np.random.Generator) -> CalibrationFit:
    motor, measured = _arrays(rows)
    home_guess = np.array([168.486, 179.824, 355.869, 82.881])
    lower = np.r_[home_guess - 180.0, [-500.0, -500.0, -500.0, -180.0]]
    upper = np.r_[home_guess + 180.0, [500.0, 500.0, 500.0, 180.0]]
    best = None
    for start_index in range(starts):
        if start_index == 0:
            x0 = np.r_[home_guess, [0.0, 0.0, 0.0, 0.0]]
        else:
            x0 = np.r_[home_guess + rng.normal(0.0, 45.0, 4),
                       rng.normal(0.0, 100.0, 3), rng.uniform(-90.0, 90.0)]
            x0 = np.clip(x0, lower + 1e-6, upper - 1e-6)
        result = least_squares(
            calibration_residual, x0, args=(signs, motor, measured),
            bounds=(lower, upper), loss=ROBUST_LOSS,
            f_scale=ROBUST_F_SCALE_MM, max_nfev=5000,
        )
        if best is None or result.cost < best.cost:
            best = result
    assert best is not None
    return CalibrationFit(
        signs=signs.copy(), zeros_deg=best.x[:4].copy(),
        translation_mm=best.x[4:7].copy(), yaw_deg=float(best.x[7]),
        robust_cost=float(best.cost), residual_vector=best.fun.copy(),
        success=bool(best.success),
    )


def fit_calibration(rows: Sequence[dict], starts: int = CALIBRATION_MULTISTARTS,
                    seed: int = 20260907) -> tuple[CalibrationFit, list[CalibrationFit]]:
    rng = np.random.default_rng(seed)
    fits = [_fit_one_sign(np.array(signs, dtype=float), rows, starts, rng)
            for signs in product((1.0, -1.0), repeat=4)]
    fits.sort(key=lambda fit: fit.robust_cost)
    return fits[0], fits


def run_fk_validation(rows: Sequence[dict], fit: CalibrationFit) -> pd.DataFrame:
    records = []
    for row in rows:
        motors = np.array([row[f"MotorDeg{i}"] for i in range(11, 15)])
        q = motor_deg_to_q(motors, fit.signs, fit.zeros_deg)
        predicted = forward_kinematics_work(q, fit)
        measured = np.array([row["Measured_X_mm"], row["Measured_Y_mm"], row["Measured_Z_mm"]])
        error = predicted - measured
        records.append({
            "Pose": row["Pose"], "q1_deg": q[0], "q2_deg": q[1], "q3_deg": q[2], "q4_deg": q[3],
            "Tool_Pitch_deg": tool_pitch_deg(q),
            "FK_X_mm": predicted[0], "FK_Y_mm": predicted[1], "FK_Z_mm": predicted[2],
            "Measured_X_mm": measured[0], "Measured_Y_mm": measured[1], "Measured_Z_mm": measured[2],
            "Error_X_mm": error[0], "Error_Y_mm": error[1], "Error_Z_mm": error[2],
            "Total_Error_mm": float(np.linalg.norm(error)),
        })
    return pd.DataFrame(records)


def run_leave_one_out(rows: Sequence[dict]) -> pd.DataFrame:
    results = []
    for held_index, held in enumerate(rows):
        training = [row for i, row in enumerate(rows) if i != held_index]
        fit, ranking = fit_calibration(training, starts=LOO_MULTISTARTS, seed=20260907 + held_index)
        motors = np.array([held[f"MotorDeg{i}"] for i in range(11, 15)])
        q = motor_deg_to_q(motors, fit.signs, fit.zeros_deg)
        predicted = forward_kinematics_work(q, fit)
        measured = np.array([held["Measured_X_mm"], held["Measured_Y_mm"], held["Measured_Z_mm"]])
        error = predicted - measured
        gap = ranking[1].robust_cost - ranking[0].robust_cost
        results.append({
            "Pose": held["Pose"], "Holdout_Error_mm": float(np.linalg.norm(error)),
            "Error_X_mm": error[0], "Error_Y_mm": error[1], "Error_Z_mm": error[2],
            "Best_Signs": " ".join(f"{int(s):+d}" for s in fit.signs),
            "Runner_Up_Cost_Gap": gap,
        })
    return pd.DataFrame(results)


def numerical_jacobian(q_rad: np.ndarray) -> np.ndarray:
    jacobian = np.zeros((3, 4))
    for joint in range(4):
        plus = q_rad.copy(); plus[joint] += IK_FINITE_DIFFERENCE_RAD
        minus = q_rad.copy(); minus[joint] -= IK_FINITE_DIFFERENCE_RAD
        jacobian[:, joint] = (
            forward_kinematics_robot(np.degrees(plus)) - forward_kinematics_robot(np.degrees(minus))
        ) / (2.0 * IK_FINITE_DIFFERENCE_RAD)
    return jacobian


def solve_ik_dls(target_work_mm: np.ndarray, fit: CalibrationFit,
                 seeds_deg: Iterable[Sequence[float]], preferred_pitch_deg: float | None = None) -> IKResult:
    target_robot = work_to_robot(target_work_mm, fit.translation_mm, fit.yaw_deg)
    lower = np.radians(JOINT_LIMITS_DEG[:, 0])
    upper = np.radians(JOINT_LIMITS_DEG[:, 1])
    candidates: list[IKResult] = []

    for seed in seeds_deg:
        q = np.clip(np.radians(np.asarray(seed, dtype=float)), lower, upper)
        reason = "MAX_ITERATIONS"
        for iteration in range(1, IK_MAX_ITERATIONS + 1):
            current_robot = forward_kinematics_robot(np.degrees(q))
            error = target_robot - current_robot
            error_norm = float(np.linalg.norm(error))
            if error_norm <= IK_TOLERANCE_MM:
                reconstructed = forward_kinematics_work(np.degrees(q), fit)
                work_error = reconstructed - target_work_mm
                candidates.append(IKResult(True, np.degrees(q), reconstructed, work_error,
                                           float(np.linalg.norm(work_error)), iteration, "CONVERGED"))
                break
            jacobian = numerical_jacobian(q)
            system = jacobian @ jacobian.T + (IK_DAMPING ** 2) * np.eye(3)
            try:
                dq = jacobian.T @ np.linalg.solve(system, error)
            except np.linalg.LinAlgError:
                reason = "SINGULAR_NUMERICAL_SYSTEM"
                break
            max_step = math.radians(IK_MAX_STEP_DEG)
            dq = np.clip(dq, -max_step, max_step)
            next_q = np.clip(q + dq, lower, upper)
            if np.linalg.norm(next_q - q) < 1e-10:
                reason = "STALLED_AT_JOINT_LIMIT"
                break
            q = next_q
        else:
            iteration = IK_MAX_ITERATIONS
        if not any(c.converged and np.allclose(c.q_deg, np.degrees(q), atol=1e-8) for c in candidates):
            reconstructed = forward_kinematics_work(np.degrees(q), fit)
            work_error = reconstructed - target_work_mm
            candidates.append(IKResult(False, np.degrees(q), reconstructed, work_error,
                                       float(np.linalg.norm(work_error)), iteration, reason))

    def score(result: IKResult) -> tuple[float, float, float]:
        pitch_penalty = 0.0 if preferred_pitch_deg is None else abs(tool_pitch_deg(result.q_deg) - preferred_pitch_deg)
        posture_penalty = float(np.linalg.norm(result.q_deg))
        return (0.0 if result.converged else 1.0, result.error_mm + 0.02 * pitch_penalty, posture_penalty)

    return min(candidates, key=score)


def run_ik_validation(rows: Sequence[dict], fit: CalibrationFit) -> pd.DataFrame:
    base_seeds = [
        [0.0, 0.0, 0.0, 0.0], [0.0, -45.0, 45.0, 0.0],
        [0.0, 45.0, -45.0, 0.0], [90.0, 0.0, 0.0, 0.0],
        [-90.0, 0.0, 0.0, 0.0],
    ]
    previous: np.ndarray | None = None
    records = []
    for row in rows:
        target = np.array([row["Measured_X_mm"], row["Measured_Y_mm"], row["Measured_Z_mm"]])
        seeds = ([previous] if previous is not None else []) + base_seeds
        result = solve_ik_dls(target, fit, seeds)
        if result.converged:
            previous = result.q_deg.copy()
        motor_deg = q_to_motor_deg(result.q_deg, fit.signs, fit.zeros_deg)
        raw = motor_deg_to_raw(motor_deg)
        records.append({
            "Pose": row["Pose"], "Target_X_mm": target[0], "Target_Y_mm": target[1], "Target_Z_mm": target[2],
            "IK_q1_deg": result.q_deg[0], "IK_q2_deg": result.q_deg[1],
            "IK_q3_deg": result.q_deg[2], "IK_q4_deg": result.q_deg[3],
            "IK_Tool_Pitch_deg": tool_pitch_deg(result.q_deg),
            **{f"IK_MotorDeg{i+11}": motor_deg[i] for i in range(4)},
            **{f"IK_RAW{i+11}": raw[i] for i in range(4)},
            "Reconstructed_X_mm": result.reconstructed_work_mm[0],
            "Reconstructed_Y_mm": result.reconstructed_work_mm[1],
            "Reconstructed_Z_mm": result.reconstructed_work_mm[2],
            "Error_X_mm": result.error_vector_mm[0], "Error_Y_mm": result.error_vector_mm[1],
            "Error_Z_mm": result.error_vector_mm[2], "Total_Error_mm": result.error_mm,
            "IK_Converged": result.converged, "IK_Iterations": result.iterations,
            "IK_Status": result.reason,
        })
    return pd.DataFrame(records)


def _style_workbook(path: Path) -> None:
    from openpyxl import load_workbook
    from openpyxl.styles import Alignment, Font, PatternFill

    wb = load_workbook(path)
    header_fill = PatternFill("solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True, name="Arial", size=10)
    body_font = Font(name="Arial", size=10)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = ws.dimensions
        ws.sheet_view.showGridLines = False
        for cell in ws[1]:
            cell.fill = header_fill; cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        for row in ws.iter_rows(min_row=2):
            for cell in row:
                cell.font = body_font
                if isinstance(cell.value, float): cell.number_format = "0.000"
        for column in ws.columns:
            letter = column[0].column_letter
            width = max(len(str(c.value)) if c.value is not None else 0 for c in column) + 2
            ws.column_dimensions[letter].width = min(max(width, 11), 24)
        ws.row_dimensions[1].height = 34
    wb.save(path)


def save_excel(path: Path, input_df: pd.DataFrame, fk_df: pd.DataFrame,
               ik_df: pd.DataFrame, fit: CalibrationFit, ranking: list[CalibrationFit],
               loo_df: pd.DataFrame) -> None:
    errors = fk_df["Total_Error_mm"].to_numpy()
    ranking_gap = ranking[1].robust_cost - ranking[0].robust_cost
    parameters = [
        ("Best fit", "sign1", int(fit.signs[0])), ("Best fit", "sign2", int(fit.signs[1])),
        ("Best fit", "sign3", int(fit.signs[2])), ("Best fit", "sign4", int(fit.signs[3])),
        *[("Best fit", f"zero_deg{i+1}", fit.zeros_deg[i]) for i in range(4)],
        ("Best fit", "tx_mm", fit.translation_mm[0]), ("Best fit", "ty_mm", fit.translation_mm[1]),
        ("Best fit", "tz_mm", fit.translation_mm[2]), ("Best fit", "yaw_offset_deg", fit.yaw_deg),
        ("Physical FK", "mean_error_mm", float(np.mean(errors))),
        ("Physical FK", "rms_error_mm", float(np.sqrt(np.mean(errors ** 2)))),
        ("Physical FK", "max_error_mm", float(np.max(errors))),
        ("Physical FK", "worst_pose", fk_df.loc[fk_df["Total_Error_mm"].idxmax(), "Pose"]),
        ("Identifiability", "runner_up_robust_cost_gap", ranking_gap),
        ("Identifiability", "warning", "q1 zero and work-frame yaw can be strongly coupled; inspect sign ranking"),
    ]
    for _, row in loo_df.iterrows():
        parameters.append(("Leave one out", f"{row['Pose']} holdout_error_mm", row["Holdout_Error_mm"]))
    for rank, candidate in enumerate(ranking, 1):
        parameters.append(("Sign ranking", f"rank_{rank}_signs", " ".join(f"{int(s):+d}" for s in candidate.signs)))
        parameters.append(("Sign ranking", f"rank_{rank}_robust_cost", candidate.robust_cost))
    parameter_df = pd.DataFrame(parameters, columns=["Section", "Parameter", "Value"])

    path.parent.mkdir(parents=True, exist_ok=True)
    with pd.ExcelWriter(path, engine="openpyxl") as writer:
        input_df.to_excel(writer, sheet_name="Calibration_Input", index=False)
        fk_df.to_excel(writer, sheet_name="FK_Validation", index=False)
        ik_df.to_excel(writer, sheet_name="IK_FK_Validation", index=False)
        parameter_df.to_excel(writer, sheet_name="Calibration_Parameters", index=False)
        loo_df.to_excel(writer, sheet_name="Calibration_Parameters", index=False, startrow=len(parameter_df) + 3)
    _style_workbook(path)


def save_plots(output_dir: Path, fk_df: pd.DataFrame, ik_df: pd.DataFrame) -> None:
    try:
        import matplotlib.pyplot as plt
    except ImportError:
        print("Optional plots skipped: matplotlib is not installed.")
        return
    output_dir.mkdir(parents=True, exist_ok=True)
    fig = plt.figure(figsize=(8, 6)); ax = fig.add_subplot(111, projection="3d")
    ax.scatter(fk_df.Measured_X_mm, fk_df.Measured_Y_mm, fk_df.Measured_Z_mm, label="Measured")
    ax.scatter(fk_df.FK_X_mm, fk_df.FK_Y_mm, fk_df.FK_Z_mm, label="FK predicted")
    ax.set(xlabel="X mm", ylabel="Y mm", zlabel="Z mm", title="Measured and FK-predicted XYZ"); ax.legend()
    fig.tight_layout(); fig.savefig(output_dir / "fk_measured_vs_predicted.png", dpi=160); plt.close(fig)
    fig, ax = plt.subplots(figsize=(8, 4)); ax.bar(fk_df.Pose, fk_df.Total_Error_mm)
    ax.set(ylabel="Error mm", title="Physical FK error by pose"); fig.tight_layout()
    fig.savefig(output_dir / "fk_error_by_pose.png", dpi=160); plt.close(fig)
    fig, ax = plt.subplots(figsize=(8, 4)); ax.bar(ik_df.Pose, ik_df.Total_Error_mm)
    ax.set(ylabel="Error mm", title="IK to FK reconstruction error"); fig.tight_layout()
    fig.savefig(output_dir / "ik_fk_error_by_pose.png", dpi=160); plt.close(fig)


def _print_table(df: pd.DataFrame, columns: list[str]) -> None:
    print(df[columns].to_string(index=False, float_format=lambda value: f"{value:9.3f}"))


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, default=Path(__file__).resolve().parent / "outputs" / "01a07b6c-8aba-76d3-b7e9-5182967627c8" / OUTPUT_FILENAME)
    parser.add_argument("--plots", action="store_true", help="Save optional PNG plots beside the workbook")
    args = parser.parse_args()

    input_df = pd.DataFrame(POSES)
    fit, ranking = fit_calibration(POSES)
    fk_df = run_fk_validation(POSES, fit)
    loo_df = run_leave_one_out(POSES)
    ik_df = run_ik_validation(POSES, fit)
    save_excel(args.output, input_df, fk_df, ik_df, fit, ranking, loo_df)
    if args.plots: save_plots(args.output.parent, fk_df, ik_df)

    print("\n" + "=" * 60 + "\nBEST MOTOR -> JOINT CALIBRATION\n" + "=" * 60)
    for i in range(4): print(f"J{i+1}: sign={int(fit.signs[i]):+d}, zero={fit.zeros_deg[i]:.6f} deg")
    print(f"work transform: tx={fit.translation_mm[0]:.3f}, ty={fit.translation_mm[1]:.3f}, tz={fit.translation_mm[2]:.3f} mm, yaw={fit.yaw_deg:.6f} deg")
    print(f"runner-up robust-cost gap: {ranking[1].robust_cost - ranking[0].robust_cost:.6f}")

    print("\n" + "=" * 60 + "\nFK PHYSICAL VALIDATION\n" + "=" * 60)
    _print_table(fk_df, ["Pose", "Measured_X_mm", "Measured_Y_mm", "Measured_Z_mm", "FK_X_mm", "FK_Y_mm", "FK_Z_mm", "Total_Error_mm"])
    errors = fk_df.Total_Error_mm.to_numpy()
    print(f"Mean={np.mean(errors):.3f} mm  RMS={np.sqrt(np.mean(errors**2)):.3f} mm  Max={np.max(errors):.3f} mm")
    print(f"Worst-fitting pose: {fk_df.loc[fk_df.Total_Error_mm.idxmax(), 'Pose']}")

    print("\n" + "=" * 60 + "\nLEAVE-ONE-OUT TEST\n" + "=" * 60)
    _print_table(loo_df, ["Pose", "Holdout_Error_mm", "Error_X_mm", "Error_Y_mm", "Error_Z_mm"])
    loo_worst = loo_df.loc[loo_df.Holdout_Error_mm.idxmax()]
    print(f"Largest holdout error: {loo_worst.Pose} = {loo_worst.Holdout_Error_mm:.3f} mm")

    print("\n" + "=" * 60 + "\nIK -> FK VALIDATION\n" + "=" * 60)
    _print_table(ik_df, ["Pose", "Target_X_mm", "Target_Y_mm", "Target_Z_mm", "Total_Error_mm", "IK_Iterations"])
    print("IK status:", ", ".join(f"{r.Pose}={r.IK_Status}" for _, r in ik_df.iterrows()))

    physical_plausible = (
        np.mean(errors) <= PHYSICAL_MEAN_ERROR_LIMIT_MM
        and np.max(errors) <= PHYSICAL_MAX_ERROR_LIMIT_MM
        and loo_df.Holdout_Error_mm.max() <= PHYSICAL_LOO_ERROR_LIMIT_MM
    )
    ik_passed = bool(ik_df.IK_Converged.all() and ik_df.Total_Error_mm.max() <= IK_TOLERANCE_MM)
    pose5_holdout = float(loo_df.loc[loo_df.Pose == "Pose 5", "Holdout_Error_mm"].iloc[0])
    median_loo = float(loo_df.Holdout_Error_mm.median())
    pose5_outlier = pose5_holdout > max(50.0, 2.5 * median_loo)

    print("\n" + "=" * 60 + "\nFINAL CONCLUSION\n" + "=" * 60)
    print("FK calibration looks plausible." if physical_plausible else "FK calibration is not yet physically reliable.")
    print("IK->FK mathematical reconstruction passed." if ik_passed else "IK solver still has significant errors.")
    print(f"Pose 5 holdout evidence: {pose5_holdout:.3f} mm; " + ("appears suspicious/outlying." if pose5_outlier else "is not uniquely identified as an outlier."))
    print("A good IK->FK result does not prove physical FK calibration.")
    if not physical_plausible:
        print("GO/NO-GO: NO-GO for real XYZ/IK motion. Remeasure physical TCP poses first.")
    print("No serial port was opened and no robot command was sent.")
    print(f"Workbook saved: {args.output}")


if __name__ == "__main__":
    main()
